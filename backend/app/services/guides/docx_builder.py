import io
from datetime import datetime, timezone

from docx import Document

from app.services.guides.rows import GUIDE_TABLE_HEADER, GuideRow, format_decimal

REVIEW_WARNING = (
    "Review every answer before use. This guide was transcribed automatically from a "
    "photograph and may contain errors. A wrong answer here will mark every script wrong."
)


def build_guide_docx(title: str, subject: str | None, rows: list[GuideRow]) -> bytes:
    doc = Document()
    doc.add_heading(title, level=0)
    doc.add_paragraph(f"Subject: {subject or 'not set'}")
    doc.add_paragraph(f"Generated: {datetime.now(timezone.utc).date().isoformat()}")

    doc.add_paragraph().add_run(REVIEW_WARNING).bold = True
    doc.add_paragraph(
        "Edit the table below, then upload this file to convert it to a mark scheme. "
        "Keep the header row. Marks must be a number; Matcher must be numeric or exact. "
        "Leave Sub-part empty for questions without one."
    )

    table = doc.add_table(rows=1, cols=len(GUIDE_TABLE_HEADER))
    table.style = "Table Grid"
    for cell, heading in zip(table.rows[0].cells, GUIDE_TABLE_HEADER):
        cell.paragraphs[0].add_run(heading).bold = True

    for row in rows:
        values = (
            row["question"],
            row["sub_part"] or "",
            row["answer"],
            format_decimal(row["marks"]),
            row["matcher"],
        )
        for cell, value in zip(table.add_row().cells, values):
            cell.text = value

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
