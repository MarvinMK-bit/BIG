import io
from decimal import Decimal, InvalidOperation

from docx import Document

from app.services.grading.schemes import VALID_MATCHERS
from app.services.guides.rows import GUIDE_TABLE_HEADER, GuideRow

_EXPECTED_HEADER = " | ".join(GUIDE_TABLE_HEADER)


def _parse_marks(text: str, where: str) -> Decimal:
    try:
        marks = Decimal(text)
    except InvalidOperation:
        raise ValueError(f"{where}: Marks must be a number, got {text!r}") from None
    if not marks.is_finite() or marks <= 0:
        raise ValueError(f"{where}: Marks must be a number greater than zero, got {text!r}")
    return marks


def parse_guide_docx(file_bytes: bytes) -> list[GuideRow]:
    """Read the reviewed guide table back out of a DOCX. Raises ValueError on any problem."""
    try:
        doc = Document(io.BytesIO(file_bytes))
    except Exception:
        # python-docx raises several unrelated types (zip, XML, missing part) for a bad file
        raise ValueError("This file isn't a readable Word document (.docx).") from None

    if not doc.tables:
        raise ValueError(f"No table found. The guide must contain a table with columns {_EXPECTED_HEADER}.")
    table = doc.tables[0]

    header = tuple(cell.text.strip() for cell in table.rows[0].cells) if table.rows else ()
    if tuple(h.casefold() for h in header) != tuple(h.casefold() for h in GUIDE_TABLE_HEADER):
        raise ValueError(
            f"The first table's header must be {_EXPECTED_HEADER}, got {' | '.join(header) or 'nothing'}."
        )

    rows: list[GuideRow] = []
    # Row numbers count the header as row 1, matching what the admin sees in Word
    for number, table_row in enumerate(table.rows[1:], start=2):
        cells = [cell.text.strip() for cell in table_row.cells]
        if not any(cells):
            continue  # Word users often leave a blank row at the end
        where = f"Table row {number}"
        if len(cells) != len(GUIDE_TABLE_HEADER):
            raise ValueError(f"{where} has {len(cells)} cells, expected {len(GUIDE_TABLE_HEADER)}.")

        question, sub_part, answer, marks, matcher = cells
        if not question:
            raise ValueError(f"{where}: Question is empty.")
        where = f"{where} (question {question}{f' part {sub_part}' if sub_part else ''})"
        if not answer:
            raise ValueError(f"{where}: Answer is empty.")
        matcher = matcher.casefold()
        if matcher not in VALID_MATCHERS:
            raise ValueError(
                f"{where}: Matcher must be one of {', '.join(sorted(VALID_MATCHERS))}, got {cells[4]!r}."
            )

        rows.append(
            GuideRow(
                question=question,
                sub_part=sub_part or None,
                answer=answer,
                marks=_parse_marks(marks, where),
                matcher=matcher,
            )
        )

    if not rows:
        raise ValueError("The guide table has no question rows.")
    return rows
